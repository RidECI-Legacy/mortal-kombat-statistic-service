import io
import base64
import asyncio
import pandas as pd
import matplotlib.pyplot as plt
from datetime import datetime
from xhtml2pdf import pisa
from src.rideci.core.logging import logger
from openpyxl.styles import Font, PatternFill, Border, Side
from openpyxl.drawing.image import Image as OpenPyXLImage
from src.rideci.infrastructure.storage.S3StorageService import S3StorageService
from src.rideci.domain.services.AIAnalysisService import AIAnalysisService

class ReportGeneratorUseCase:
    def __init__(self, s3_service: S3StorageService, ai_service: AIAnalysisService):
        self.s3_service = s3_service
        self.ai_service = ai_service

    async def execute(self, user_id: str, data: list, user_stats: dict, report_format: str):
        try:
            df = pd.DataFrame(data)
            badge = user_stats.get("currentBadge", "N/A").replace("BadgeType.", "").replace("_", " ").title()

            loop = asyncio.get_running_loop()
            
            
            if report_format == 'excel':
                buffer = await loop.run_in_executor(None, self._generate_excel_report, df, user_id, badge)
                file_name = f"reports/RidECI_{user_id}_{datetime.now().strftime('%Y%m%d')}.xlsx"
            else:
                buffer = await self._generate_pdf_report(df, user_id, badge)
                file_name = f"reports/RidECI_{user_id}_{datetime.now().strftime('%Y%m%d')}.pdf"
                
            return self.s3_service.upload_file(buffer, file_name)
        except Exception as e:
            logger.error(f"Error generando reporte para {user_id}: {str(e)}")
            raise

    def _generate_chart_buffer(self, df: pd.DataFrame) -> io.BytesIO:
        plt.style.use('ggplot')
        fig, ax = plt.subplots(figsize=(5, 4))
        ax.pie(df['Valor'], labels=df['Metrica'], autopct='%1.1f%%', colors=['#2E7D32', '#1565C0', '#F9A825'])
        buf = io.BytesIO()
        plt.savefig(buf, format='png', bbox_inches='tight')
        plt.close(fig)  
        buf.seek(0)
        return buf

    def _generate_excel_report(self, df: pd.DataFrame, user_id: str, badge: str) -> io.BytesIO:
        buffer = io.BytesIO()
        with pd.ExcelWriter(buffer, engine='openpyxl') as writer:
            meta = [["REPORTE RIDECI"], ["Usuario", user_id], ["Fecha", datetime.now().strftime("%Y-%m-%d %H:%M")], ["Medalla", badge]]
            pd.DataFrame(meta).to_excel(writer, index=False, header=False, sheet_name='Reporte')
            df.to_excel(writer, index=False, startrow=6, sheet_name='Reporte')
            
            ws = writer.sheets['Reporte']
            thin_border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))
            
            for row in ws.iter_rows(min_row=7, max_row=7+len(df), min_col=1, max_col=2):
                for cell in row: cell.border = thin_border
            
            header_fill = PatternFill(start_color="1B5E20", end_color="1B5E20", fill_type="solid")
            for cell in ws[7]:
                cell.fill, cell.font = header_fill, Font(color="FFFFFF", bold=True)
            
            img = OpenPyXLImage(self._generate_chart_buffer(df))
            ws.add_image(img, 'E2')
            
        buffer.seek(0)
        return buffer

    async def _generate_pdf_report(self, df: pd.DataFrame, user_id: str, badge: str) -> io.BytesIO:
        buffer = io.BytesIO()
        stats_json = df.to_json(orient='records')
        try:
            ia_html = await self.ai_service.analyze(stats_json, badge)
        except Exception as e:
            logger.warning(f"IA no disponible para {user_id}, usando plantilla estática: {str(e)}")
            ia_html = f"""
                <div style="border: 1px solid #ccc; padding: 10px; background-color: #f9f9f9;">
                    <p style="font-size: 10px; color: #666;">
                        <em>Aviso: El análisis cualitativo IA no está disponible temporalmente. 
                        Los datos presentados son estadísticos.</em>
                    </p>
                </div>
            """
        pie_b64 = base64.b64encode(self._generate_chart_buffer(df).read()).decode('utf-8')
        
        html = f"""
        <html>
            <style>
                @page {{ size: A4; margin: 1.5cm; }}
                body {{ font-family: Helvetica, sans-serif; color: #333; font-size: 12px; }}
                .logo {{ width: 40px; height: 40px; border: 1px dashed #1B5E20; text-align: center; line-height: 40px; float: left; margin-right: 10px; font-size: 8px; }}
                .header {{ border-bottom: 2px solid #1B5E20; margin-bottom: 15px; padding-bottom: 5px; }}
                table {{ width: 100%; border-collapse: collapse; margin-top: 10px; }}
                th {{ background: #1B5E20; color: white; padding: 8px; }}
                td {{ border: 1px solid #999; padding: 6px; text-align: center; }}
            </style>
            <body>
                <div class="header">
                    <div class="logo">LOGO</div>
                    <h1>Reporte RidECI</h1>
                    <p><b>Usuario:</b> {user_id} | <b>Fecha:</b> {datetime.now().strftime("%Y-%m-%d %H:%M")}</p>
                </div>
                <h3>📊 Métricas de Desempeño</h3>
                <table>
                    <tr><th>Métrica</th><th>Valor</th></tr>
                    {"".join([f"<tr><td>{r['Metrica']}</td><td>{r['Valor']}</td></tr>" for _, r in df.iterrows()])}
                </table>
                <div style="text-align:center; margin-top:15px;">
                    <img src="data:image/png;base64,{pie_b64}" width="400">
                </div>
                <div style="margin-top:15px;">{ia_html}</div>
            </body>
        </html>
        """
        pisa.CreatePDF(html, dest=buffer)
        buffer.seek(0)
        return buffer